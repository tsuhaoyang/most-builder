"""匯出整合測試：wi-preview / excel / lb-csv / lb-api。

特別守住「openpyxl 等 runtime 依賴漏宣告」這類回歸——excel 端點會載 openpyxl，
依賴缺失時會 500。需 demo worksheet（dev_seed_v2 的 5555…）；未種則 skip。
"""
from __future__ import annotations

import io
import uuid

import pytest

pytestmark = pytest.mark.integration

WS = "55555555-5555-5555-5555-555555555555"
OBJ = "66666666-6666-6666-6666-666666666666"

# ADR-011 加法原則：lb-csv 合約凍結——LB 拿 normal（second），寬放屬 MOST 端，欄位不得增刪改序。
LB_CSV_HEADER = ("seq_no,content,second,ascription,level,level_min,level_max,variable,"
                 "countersignature,parent_countersignature,order,number,number_count,machine_count,manpower")


def _gm_row():
    return {"id": str(uuid.uuid4()), "seq_no": 1, "hand": "RH", "object_vocab_id": OBJ,
            "frequency": 1, "narrative": "export allowance",
            "cycle": {"seq": "GM", "rule_set_code": "MINIMOST_FACTORY_V1",
                      "a0": {"reach_cm": 20}, "g2": {"g_code": "g_grasp"},
                      "a3": {"reach_cm": 25}, "p5": {"p_base_code": "p_place_none"}},
            "level": {"ascription": "main", "level": "1"}}


async def _seeded(client) -> bool:
    r = await client.get(f"/api/v2/worksheets/{WS}/export/wi-preview")
    return r.status_code == 200


async def test_wi_preview(client):
    if not await _seeded(client):
        pytest.skip("demo worksheet 未種（先跑 dev_seed_v2.py）")
    r = await client.get(f"/api/v2/worksheets/{WS}/export/wi-preview")
    assert r.status_code == 200
    assert "total_tmu" in r.json()


async def test_export_excel(client):
    """守 openpyxl：缺依賴會 500。"""
    if not await _seeded(client):
        pytest.skip("demo worksheet 未種")
    r = await client.get(f"/api/v2/worksheets/{WS}/export/excel")
    assert r.status_code == 200, f"excel 匯出失敗（檢查 openpyxl 是否在 runtime 依賴）：{r.status_code}"
    ct = r.headers.get("content-type", "")
    assert "spreadsheet" in ct or "octet-stream" in ct, ct


async def test_export_excel_normal_standard_seconds(client):
    """impl-02 §3：Excel 增正常秒/標準秒欄；allowance=10 → 標準秒 = 正常秒 × 1.1。"""
    if not await _seeded(client):
        pytest.skip("demo worksheet 未種")
    from openpyxl import load_workbook

    new = (await client.post(f"/api/v2/worksheets/{WS}/clone")).json()["new_worksheet_id"]
    r = await client.put(f"/api/v2/worksheets/{new}", json={"rows": [_gm_row()], "allowance_percent": 10})
    assert r.status_code == 200, r.text
    xls = await client.get(f"/api/v2/worksheets/{new}/export/excel")
    assert xls.status_code == 200
    sheet = load_workbook(io.BytesIO(xls.content)).active
    meta = {}  # 第 2 列：label-value 對（正常秒/寬放%/標準秒）
    row2 = [c.value for c in sheet[2]]
    for i in range(0, len(row2) - 1, 2):
        meta[row2[i]] = row2[i + 1]
    assert meta.get("正常秒") and meta["正常秒"] > 0
    assert meta.get("寬放%") == 10
    assert meta.get("標準秒") == pytest.approx(meta["正常秒"] * 1.1, abs=1e-4)


async def test_export_excel_no_allowance_standard_blank(client):
    """OQ-002：allowance 未設 → 標準秒留空（不得以 normal 假充 standard）。"""
    if not await _seeded(client):
        pytest.skip("demo worksheet 未種")
    from openpyxl import load_workbook

    new = (await client.post(f"/api/v2/worksheets/{WS}/clone")).json()["new_worksheet_id"]
    r = await client.put(f"/api/v2/worksheets/{new}", json={"rows": [_gm_row()], "allowance_percent": None})
    assert r.status_code == 200, r.text
    xls = await client.get(f"/api/v2/worksheets/{new}/export/excel")
    sheet = load_workbook(io.BytesIO(xls.content)).active
    row2 = [c.value for c in sheet[2]]
    meta = {row2[i]: row2[i + 1] for i in range(0, len(row2) - 1, 2)}
    assert meta.get("正常秒") and meta["正常秒"] > 0
    assert meta.get("寬放%") in (None, "")
    assert meta.get("標準秒") in (None, "")


async def test_export_lb_csv(client):
    if not await _seeded(client):
        pytest.skip("demo worksheet 未種")
    r = await client.get(f"/api/v2/worksheets/{WS}/export/lb-csv")
    assert r.status_code == 200
    assert "text/csv" in r.headers.get("content-type", "") or len(r.content) > 0


async def test_export_lb_csv_contract_frozen(client):
    """ADR-011：lb-csv 表頭凍結；設了 allowance 也不得混入 standard/allowance 欄（LB 拿 normal）。"""
    if not await _seeded(client):
        pytest.skip("demo worksheet 未種")
    new = (await client.post(f"/api/v2/worksheets/{WS}/clone")).json()["new_worksheet_id"]
    r = await client.put(f"/api/v2/worksheets/{new}", json={"rows": [_gm_row()], "allowance_percent": 10})
    assert r.status_code == 200, r.text
    csv_r = await client.get(f"/api/v2/worksheets/{new}/export/lb-csv")
    assert csv_r.status_code == 200
    header = csv_r.text.splitlines()[0].strip().lstrip("﻿")  # 端點帶 UTF-8 BOM（Excel 相容，既有行為）
    assert header == LB_CSV_HEADER, f"lb-csv 合約被改動：{header}"


async def test_export_lb_api_dry_run(client):
    if not await _seeded(client):
        pytest.skip("demo worksheet 未種")
    r = await client.post(f"/api/v2/worksheets/{WS}/export/lb-api")
    assert r.status_code == 200
    assert "payload" in r.json()


async def test_export_report_xlsx_happy_path(client):
    """三 sheet 報表：200、content-type=xlsx、Content-Disposition 含 MOST_Report。"""
    if not await _seeded(client):
        pytest.skip("demo worksheet 未種（先跑 dev_seed_v2.py）")
    r = await client.get(f"/api/v2/worksheets/{WS}/export/report.xlsx")
    assert r.status_code == 200, r.text
    ct = r.headers.get("content-type", "")
    assert "spreadsheet" in ct or "octet-stream" in ct, ct
    cd = r.headers.get("content-disposition", "")
    assert "MOST_Report" in cd, cd


async def test_export_report_xlsx_sheets_structure(client):
    """三 sheet 各自存在，Sheet1 含「廠區」row，Sheet2 含欄位標頭，Sheet3 含簽核欄位。"""
    if not await _seeded(client):
        pytest.skip("demo worksheet 未種（先跑 dev_seed_v2.py）")
    from openpyxl import load_workbook

    r = await client.get(f"/api/v2/worksheets/{WS}/export/report.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    # Sheet 1：案件資訊
    assert "案件資訊" in wb.sheetnames
    ws1 = wb["案件資訊"]
    keys = [ws1.cell(row=i, column=1).value for i in range(1, ws1.max_row + 1)]
    assert "廠區" in keys
    assert "總 TMU" in keys
    assert "總標準秒" in keys
    # Sheet 2：動作明細
    assert "動作明細" in wb.sheetnames
    ws2 = wb["動作明細"]
    header2 = [ws2.cell(row=1, column=c).value for c in range(1, 10)]
    assert "序號" in header2
    assert "步驟TMU" in header2
    # Sheet 3：簽核歷程
    assert "簽核歷程" in wb.sheetnames
    ws3 = wb["簽核歷程"]
    header3 = [ws3.cell(row=1, column=c).value for c in range(1, 7)]
    assert "時間" in header3
    assert "執行者" in header3


async def test_export_report_xlsx_not_found(client):
    """不存在的 worksheet_id → 404。"""
    fake_id = uuid.uuid4()
    r = await client.get(f"/api/v2/worksheets/{fake_id}/export/report.xlsx")
    assert r.status_code == 404
