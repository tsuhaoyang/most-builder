"""Excel 匯入服務（Phase 2a/2b，ADR-013）：解析 → 暫存 → 欄位對應 → 正規化預覽 → 提交工序表。
Phase 2b：submit_to_worksheet — staged rows → WiRow + MostCycle(stub) + LevelEntry。

容錯：NLP 解析失敗 warning 不中斷；壞 frequency 自動置 1；DB 錯誤整批 rollback（原子性）。
MOST 仍權威（2b submit 建 stub cycle，IE 後補 MOST 分析）。
"""
from __future__ import annotations

import io
import logging
import uuid
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from uuid import UUID

    from sqlalchemy.ext.asyncio import AsyncSession

    from ddm_v2.models.v2.vocab import WorkVocabItem

from ddm_v2.schemas.v2.import_excel import REQUIRED_FIELDS
from ddm_v2.services.v2.template_matching import score_keywords

logger = logging.getLogger(__name__)

_MAX_ROWS = 1000   # 每分頁存進 staging 的上限（避免巨大 JSONB）
_MAX_COLS = 60
_HAND_MAP = {"lh": "LH", "left": "LH", "左": "LH", "左手": "LH",
             "rh": "RH", "right": "RH", "右": "RH", "右手": "RH",
             "bh": "BH", "both": "BH", "雙": "BH", "雙手": "BH"}


def parse_workbook(content: bytes) -> dict[str, Any]:
    """讀 .xlsx 全分頁 → raw_payload（截斷）。"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        grid: list[list[Any]] = []
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            if r >= _MAX_ROWS:
                break
            cells = list(row[:_MAX_COLS])
            grid.append([_cell(c) for c in cells])
        n_cols = max((len(r) for r in grid), default=0)
        sheets.append({"name": ws.title, "grid": grid, "n_rows": len(grid), "n_cols": n_cols})
    wb.close()
    return {"sheets": sheets}


def _cell(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, (int, float, str, bool)):
        return v
    return str(v)  # date/time 等轉字串


def suggest_header_row(grid: list[list[Any]]) -> int:
    """啟發式：取「非空字串儲存格最多」的列當表頭（前 15 列內）。"""
    best, best_score = 0, -1
    for i, row in enumerate(grid[:15]):
        score = sum(1 for c in row if isinstance(c, str) and c.strip())
        if score > best_score:
            best, best_score = i, score
    return best


def apply_mapping(raw_payload: dict, sheet: str, header_row: int,
                  column_map: dict[str, int], time_unit: str) -> tuple[list[dict], list[str]]:
    """套欄位對應 → 正規化列 + warnings。"""
    sheets = {s["name"]: s for s in raw_payload.get("sheets", [])}
    if sheet not in sheets:
        return [], [f"找不到分頁：{sheet}"]
    grid = sheets[sheet]["grid"]
    warnings: list[str] = []
    rows: list[dict] = []
    factor = 60.0 if time_unit == "min" else 1.0

    for ri in range(header_row + 1, len(grid)):
        raw = grid[ri]
        rec: dict[str, Any] = {"_row": ri + 1}
        for field, col in column_map.items():
            rec[field] = raw[col] if (isinstance(col, int) and 0 <= col < len(raw)) else None
        # 整列皆空 → 跳過
        if all(rec.get(f) in (None, "") for f in column_map):
            continue
        # 正規化
        if "description" in rec and rec["description"] is not None:
            rec["description"] = str(rec["description"]).strip()
        if "hand" in rec and rec.get("hand") is not None:
            rec["hand"] = _HAND_MAP.get(str(rec["hand"]).strip().lower(), None)
        if "seconds" in column_map:
            rec["seconds"], w = _num(rec.get("seconds"))
            if rec["seconds"] is not None:
                rec["seconds"] = round(rec["seconds"] * factor, 4)
            elif w:
                warnings.append(f"列{ri + 1}：工時無法解析（{raw[column_map['seconds']]!r}）")
        if "quantity" in column_map:
            rec["quantity"], _ = _num(rec.get("quantity"))
        # 必填檢查
        for rf in REQUIRED_FIELDS:
            if not rec.get(rf):
                warnings.append(f"列{ri + 1}：缺必填「{rf}」")
        rows.append(rec)
    return rows, warnings


def _num(v: Any) -> tuple[float | None, bool]:
    """轉數字；回 (值, 是否原本有值但解析失敗)。"""
    if v is None or v == "":
        return None, False
    if isinstance(v, (int, float)):
        return float(v), False
    try:
        return float(str(v).strip().replace(",", "")), False
    except ValueError:
        return None, True


# ── ADR-025 D10：匯入預覽的逐行範本建議（match + 以 active 重算 TMU）────────────────

async def build_row_matches(session: "AsyncSession", rows: list[dict]) -> list[dict]:
    """為每一暫存列附上 `match`（ADR-025 D10）。回傳「加了 match 的新列」，不改動傳入物件。

    TMU 一律走唯一計算引擎（`most_engine.compute_cycle`）以 **active** rule-set 重算——
    範本的 `cycle_template` 不含值也不含 rule_set_code（D9 已清），套用＝現在算，不是回放。

    批次策略（ADR-025「代價」節）：standard 範本查一次、active 解析一次、rule-set 載入一次；
    每個範本的 TMU **只算一次**（M 個範本），逐列只做便宜的關鍵字比對（O(N×M) 字串），
    不是每列都 load rule-set 或每列都算 TMU。N 很大時成本為 O(M) 引擎呼叫 + O(N×M) 字串比對。

    取不到 active（或 active 不完整）＝錯誤狀態：命中列的 `computed_tmu` 為 null 並帶
    `error`，**不猜版本、不 fallback**（守則 §7 第 1、8 條）。
    """
    from sqlalchemy import select

    from ddm_v2.models.v2.motion_template import MotionTemplate
    from ddm_v2.most_engine import SequenceError, compute_cycle, load_rule_set_from_db
    from ddm_v2.most_engine.rule_set_data import RuleSetIncomplete
    from ddm_v2.schemas.v2.most import CycleIn, cycle_in_to_engine
    from ddm_v2.services.v2.rule_set_service import (
        NoActiveRuleSet,
        get_active_rule_set,
    )

    # 1. standard 庫一次載入（與 /match 一致：is_active AND status='standard'）
    templates = list((await session.execute(select(MotionTemplate).where(
        MotionTemplate.is_active.is_(True), MotionTemplate.status == "standard"))).scalars().all())

    # 2. active 解析一次 + rule-set 載入一次（引擎重用，不逐列重載）
    rsdata = None
    active_error: str | None = None
    try:
        active_rs = await get_active_rule_set(session)
        rsdata = await load_rule_set_from_db(session, active_rs.code)
        rsdata.validate_complete()
    except (NoActiveRuleSet, RuleSetIncomplete) as e:
        active_error = str(e)
        rsdata = None

    # 3. 每個範本的 TMU 只算一次（命中的列與其 candidates 共用）
    computed: dict[str, dict] = {}
    for t in templates:
        entry: dict = {"computed_tmu": None, "computed_seconds": None, "error": None}
        if rsdata is None:
            entry["error"] = active_error
        else:
            try:
                cin = CycleIn.model_validate(t.cycle_template)  # cycle_template 不含 rule_set_code
                result = compute_cycle(cycle_in_to_engine(cin), rsdata)
                entry["computed_tmu"] = result.total_tmu
                entry["computed_seconds"] = result.total_seconds
            except (SequenceError, ValueError) as e:
                entry["error"] = f"範本計算失敗：{e}"
        computed[str(t.id)] = entry

    def _hit(score: float, hits: list[str], t: "MotionTemplate") -> dict:
        c = computed[str(t.id)]
        return {
            "template_id": str(t.id),
            "template_name_zh": t.name_zh,
            "seq_kind": t.seq_kind,
            "score": score,
            "matched_keywords": hits,
            "computed_tmu": c["computed_tmu"],
            "computed_seconds": c["computed_seconds"],
            "error": c["error"],
        }

    # 4. 逐列比對（便宜的字串運算）
    out: list[dict] = []
    for row in rows:
        new_row = dict(row)  # 不改動傳入的 staged 列
        desc = str(row.get("description") or "")
        scored = []
        for t in templates:
            s, hits = score_keywords(desc, list(t.keywords or []))
            if s > 0:
                scored.append((s, hits, t))
        scored.sort(key=lambda x: x[0], reverse=True)
        if not scored:
            new_row["match"] = None
        else:
            best = _hit(*scored[0])
            best["candidates"] = [_hit(*x) for x in scored[1:]]
            new_row["match"] = best
        out.append(new_row)
    return out


# ── Phase 2b：staged rows → worksheet ──────────────────────────────────────────

async def submit_to_worksheet(
    session: "AsyncSession",
    import_id: "UUID",
    worksheet_id: "UUID",
    rule_set_code: str | None,
    actor_no: str,
    row_adoptions: list | None = None,
    base_revision: int | None = None,
) -> dict:
    """Phase 2b：staged rows → WiRow + MostCycle(stub) + LevelEntry。"""
    from sqlalchemy import func, select

    from ddm_v2.models.v2.import_staging import ExcelImport
    from ddm_v2.models.v2.worksheet import LevelEntry, MostCycle, MostWorksheet, WiRow
    from ddm_v2.nlp.rule_based import RuleBasedParser
    from ddm_v2.schemas.v2.most import CycleIn
    from ddm_v2.services.v2.worksheet_revision import bump_worksheet_revision

    # ── 1. 載入 ExcelImport ──
    rec = await session.get(ExcelImport, import_id)
    if rec is None:
        raise ValueError("import_not_found")
    # Fix-H2：已提交 → 409 already_submitted；未對應 → 409 import_not_mapped（分離錯誤碼）
    if rec.status == "submitted":
        raise ValueError("already_submitted")
    if rec.status != "mapped":
        raise ValueError("import_not_mapped")
    staged = rec.staged_rows or []
    if not staged:
        raise ValueError("no_staged_rows")

    # ── 2. 載入目標 worksheet ──
    ws = await session.get(MostWorksheet, worksheet_id)
    if ws is None:
        raise ValueError("worksheet_not_found")
    # Fix-H3：只允許提交到 draft 工序表（published/retired 拒收）
    if ws.status != "draft":
        raise ValueError("worksheet_not_draft")

    # R1：內容 append 前 CAS revision
    await bump_worksheet_revision(
        session,
        worksheet_id=worksheet_id,
        base_revision=base_revision,
        edited_by=actor_no,
    )

    # ── 3. 決定 rule_set ──
    if rule_set_code:
        from ddm_v2.models.v2.rule_set import RuleSet as RuleSetModel
        rs_row = (await session.execute(
            select(RuleSetModel).where(RuleSetModel.code == rule_set_code)
        )).scalar_one_or_none()
    else:
        rs_row = None
        if ws.default_rule_set_id:
            from ddm_v2.models.v2.rule_set import RuleSet as RuleSetModel
            rs_row = await session.get(RuleSetModel, ws.default_rule_set_id)
    if rs_row is None:
        # ADR-023 §3.5：fallback ＝ **目前 active**，與 catalog/worksheet 同一條路。
        # 舊版取「最新建立的 rule_set」會抓到剛 clone 出來的未發布 draft
        # （D1 的 clone-draft 自動命名讓這條路更容易被踩），屬靜默降級。
        from ddm_v2.services.v2.rule_set_service import get_active_rule_set
        rs_row = await get_active_rule_set(session)

    # ── 4. 目前 worksheet 最大 seq_no ──
    max_seq_result = (await session.execute(
        select(func.max(WiRow.seq_no)).where(WiRow.worksheet_id == worksheet_id)
    )).scalar()
    next_seq = (max_seq_result or 0) + 1

    # ── 5. NLP parser（無 DB 同義詞，只靠內建觸發詞）──
    parser = RuleBasedParser([])

    # ── 5b. 範本採用（ADR-025 D10）──────────────────────────────────────────────
    # 只有列在 row_adoptions 的行套範本 cycle；後端重新以 active 解析並計算，
    # **不信任前端傳來的任何 TMU**（契約只收 template_id，值一律後端算）。
    # 非 standard／不存在 template_id → 422；row_index 越界 → 422（不得靜默略過）。
    adopted: dict[int, Any] = {}
    active_rs = None
    active_rsdata = None
    if row_adoptions:
        from ddm_v2.models.v2.motion_template import MotionTemplate
        from ddm_v2.most_engine import compute_cycle, load_rule_set_from_db
        from ddm_v2.schemas.v2.most import cycle_in_to_engine, resolve_cycle_rule_set
        from ddm_v2.services.v2.rule_set_service import get_active_rule_set

        ids = [ra.template_id for ra in row_adoptions]
        tmap = {t.id: t for t in (await session.execute(
            select(MotionTemplate).where(MotionTemplate.id.in_(ids)))).scalars().all()}
        for ra in row_adoptions:
            if not (0 <= ra.row_index < len(staged)):
                raise ValueError(
                    f"row_adoptions：row_index {ra.row_index} 超出暫存列範圍（可用 0..{len(staged) - 1}）")
            t = tmap.get(ra.template_id)
            if t is None or t.status != "standard" or not t.is_active:
                raise ValueError(
                    f"row_adoptions：範本 {ra.template_id} 非啟用中的標準範本"
                    "（不存在／草稿／已停用），不得套用")
            adopted[ra.row_index] = t
        # 採用行一律以 **active** 計算（範本＝新建模，不是回放）；active 載入一次、引擎重用。
        active_rs = await get_active_rule_set(session)
        active_rsdata = await load_rule_set_from_db(session, active_rs.code)
        active_rsdata.validate_complete()

    # ── 6. 逐列建立 ──
    n_with_analysis = 0
    n_need_review = 0
    warnings: list[str] = []

    for i, row in enumerate(staged):
        desc = (row.get("description") or "").strip()
        seconds = row.get("seconds")        # float | None（apply_mapping 已換算）
        hand_raw = row.get("hand")          # 'LH'|'RH'|'BH'|None
        hand = hand_raw if hand_raw in ("LH", "RH", "BH") else None

        # Fix-M5：負數或零次數改為 1（quantity 欄可能包含無效值）
        raw_qty = row.get("quantity")
        frequency = float(raw_qty or 1)
        if frequency <= 0:
            warnings.append(f"列{row.get('_row', '?')}：次數 {raw_qty!r} 非正數，改為 1")
            frequency = 1.0

        # 6a. vocab 查找或建立
        vocab = await _lookup_or_create_vocab(session, desc or "（匯入）")

        tmpl = adopted.get(i)
        if tmpl is not None:
            # adopted 非空 ⇒ 5b 已解析 active（不變式；narrow Optional）
            assert active_rs is not None and active_rsdata is not None
            # 6b-採用. 套範本 cycle，後端以 **active** 重算（ADR-025 §3；回放鐵則 §3.4：
            # 落地的 cycle 記錄「當下 active」的 code，日後回同版回放）。
            cycle_in = CycleIn.model_validate(tmpl.cycle_template)  # cycle_template 不含 rule_set_code
            resolve_cycle_rule_set(cycle_in, active_rs.code)         # 快照當下 active code
            result = compute_cycle(cycle_in_to_engine(cycle_in), active_rsdata)
            seq_kind = result.seq
            slot_inputs = cycle_in.model_dump(mode="json")
            cyc_rule_set_id = active_rs.id
            cyc_total_tmu = result.total_tmu
            cyc_total_seconds = result.total_seconds
            n_with_analysis += 1
        else:
            # 6b-stub. nl-draft 推斷 seq_kind（Fix-M1：失敗記 warning，不中斷）
            seq_kind = "GM"
            if desc:
                try:
                    nl = parser.parse(desc)
                    if nl.suggested_seq:
                        seq_kind = nl.suggested_seq
                        n_with_analysis += 1
                except Exception as exc:  # noqa: BLE001 — NLP 為可選增益，失敗不應中斷匯入
                    logger.warning("nl-draft parse failed for row %s: %s", row.get("_row"), exc)

            # 6c. stub CycleIn（IE 後續補完 MOST 分析）
            cycle_in = CycleIn(seq=seq_kind, rule_set_code=rs_row.code)
            slot_inputs = cycle_in.model_dump(mode="json")
            cyc_rule_set_id = rs_row.id
            cyc_total_tmu = 0.0
            cyc_total_seconds = 0.0  # stub：觀測工時只存 LevelEntry.raw_seconds

            if seconds is None:
                n_need_review += 1
                if desc:
                    warnings.append(f"列{row.get('_row','?')}「{desc[:20]}」無工時，需人工補 MOST")

        # 6d. WiRow
        wi = WiRow(
            id=uuid.uuid4(),
            worksheet_id=worksheet_id,
            seq_no=next_seq + i,
            object_vocab_id=vocab.id,
            hand=hand,
            frequency=frequency,
            provenance="imported",
            source_import_id=import_id,
        )
        session.add(wi)
        await session.flush()  # 取 wi.id

        # 6e. MostCycle（採用行＝實 cycle；未採用＝stub。觀測工時一律存 LevelEntry.raw_seconds）
        session.add(MostCycle(
            id=uuid.uuid4(),
            wi_row_id=wi.id,
            seq_kind=seq_kind,
            rule_set_id=cyc_rule_set_id,
            slot_inputs=slot_inputs,
            total_tmu=cyc_total_tmu,
            total_seconds=cyc_total_seconds,
        ))

        # 6f. LevelEntry（raw_seconds = 匯入觀測工時）
        session.add(LevelEntry(
            id=uuid.uuid4(),
            wi_row_id=wi.id,
            worksheet_id=worksheet_id,
            raw_seconds=round(seconds or 0.0, 4),
        ))

    # ── 7. 更新 ExcelImport 狀態 ──
    rec.status = "submitted"
    rec.submitted_worksheet_id = worksheet_id
    await session.flush()

    from ddm_v2.services.v2 import worksheet_service as ws_svc
    from ddm_v2.services.v2.worksheet_revision import content_hash_from_read, set_content_hash

    snap = await ws_svc.read_worksheet(session, worksheet_id)
    ch = content_hash_from_read(snap)
    await set_content_hash(session, worksheet_id=worksheet_id, content_hash=ch)
    await session.flush()

    return {
        "worksheet_id": str(worksheet_id),
        "n_rows": len(staged),
        "n_with_analysis": n_with_analysis,
        "n_need_review": n_need_review,
        "warnings": warnings,
        "revision_no": int(snap["revision_no"]),
        "content_hash": ch,
    }


async def _lookup_or_create_vocab(session: "AsyncSession", name_zh: str) -> "WorkVocabItem":
    """依 name_zh 查找現有 vocab，找不到就建立 imported 類型的。"""
    from sqlalchemy import func, select

    from ddm_v2.models.v2.vocab import WorkVocabItem

    existing = (await session.execute(
        select(WorkVocabItem).where(
            func.lower(WorkVocabItem.name_zh) == func.lower(name_zh.strip()),
            WorkVocabItem.is_active.is_(True),
        ).limit(1)
    )).scalar_one_or_none()
    if existing:
        return existing

    new_v = WorkVocabItem(
        id=uuid.uuid4(),
        kind="object",
        name_zh=name_zh.strip()[:200],
        source_system="imported",
    )
    session.add(new_v)
    await session.flush()
    return new_v
