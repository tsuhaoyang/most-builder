"""匯出服務：WI 1128 預覽 / Excel / LB csv / LB API 接口（payload）/ 三 sheet 管理報表。

依據 system-architecture-v2 §8.1。皆由 worksheet_service.read_worksheet 的權威資料衍生。
LB API 的 request model 由 User 後續提供；此處先做 adapter 接口（先回 dry-run payload）。
"""
from __future__ import annotations

import csv
import io
import uuid
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from ddm_v2.models.v2.audit import WorkflowAuditLog
from ddm_v2.models.v2.org import Product, Site, Sku
from ddm_v2.models.v2.worksheet import MostWorksheet, ProcessVersion, WiRow
from ddm_v2.most_engine import level as level_engine
from ddm_v2.most_engine.rule_set_data import TMU_TO_SEC
from ddm_v2.services.v2.worksheet_service import WorksheetNotFound, read_worksheet


def _parse_tech(tech: str | None) -> list[dict[str, Any]]:
    """'A6 B0 G6 A10 B0 P6 A0' → [{letter:'A',value:'6'}, ...]。"""
    out: list[dict[str, Any]] = []
    for tok in (tech or "").split():
        i = 0
        while i < len(tok) and tok[i].isalpha():
            i += 1
        out.append({"letter": tok[:i], "value": tok[i:]})
    return out


async def wi_preview(session: AsyncSession, ws_id: uuid.UUID) -> dict[str, Any]:
    data = await read_worksheet(session, ws_id)
    rows = []
    for r in data["rows"]:
        cyc = r.get("cycle") or {}
        lv = r.get("level") or {}
        rows.append({
            "seq_no": r["seq_no"], "sub_activity": r.get("sub_activity") or "",
            "key_parts": r.get("key_parts") or "", "hand": r.get("hand") or "",
            "method": cyc.get("narrative") or "", "slots": _parse_tech(cyc.get("tech_line")),
            "freq": r["frequency"], "simo": bool(r["simo_group_id"]), "tmu": cyc.get("total_tmu"),
            "level": {  # Level System 關係欄（人看得懂工序限制）
                "ascription": lv.get("ascription"), "level": lv.get("level"),
                "countersignature": lv.get("countersignature"),
                "parent_countersignature": lv.get("parent_countersignature"),
                "order": lv.get("order"), "number": lv.get("number"), "number_count": lv.get("number_count"),
            },
        })
    return {"worksheet_id": str(ws_id), "status": data["status"], "rows": rows, "total_tmu": data["total_tmu"],
            # 時間投影：normal=引擎輸出；standard=normal×(1+allowance%/100)，allowance 未設＝None
            "normal_seconds": data["normal_seconds"], "allowance_percent": data["allowance_percent"],
            "standard_seconds": data["standard_seconds"]}


async def to_excel_bytes(session: AsyncSession, ws_id: uuid.UUID) -> bytes:
    """1128 式 WI 工序單 .xlsx。"""
    from openpyxl import Workbook

    prev = await wi_preview(session, ws_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "WI"
    ws.append(["MODEL", "", "ANALYST", "", "TOTAL TMU", prev["total_tmu"]])
    # 時間欄（impl-02 §3）：正常秒＝引擎輸出；標準秒＝normal×(1+allowance%/100)。
    # allowance 未設（NULL）→ 寬放%/標準秒留空，不得以 normal 假充 standard。
    ws.append(["正常秒", prev["normal_seconds"],
               "寬放%", prev["allowance_percent"] if prev["allowance_percent"] is not None else "",
               "標準秒", prev["standard_seconds"] if prev["standard_seconds"] is not None else ""])
    ws.append([])
    header = ["STEP", "SUB活動", "Key Parts", "HAND", "METHOD",
              "A", "B", "G", "A/M", "B/X", "P/I", "A", "Freq", "SIMO", "TMU",
              # ── Level System：工序關係（給 LB / 給人看）──
              "主序Level", "角色", "群組", "巢狀母組", "組內序", "不可同站nb", "nb上限"]
    ws.append(header)
    for r in prev["rows"]:
        slots = r["slots"] + [{"value": ""}] * (7 - len(r["slots"]))
        vals = [s["value"] for s in slots[:7]]
        lv = r.get("level") or {}
        cs = lv.get("countersignature") or ""
        role = "主" if (lv.get("ascription") == "main") else ("屬" if cs else "")
        ws.append([r["seq_no"], r["sub_activity"], r["key_parts"], r["hand"], r["method"],
                   *vals, r["freq"], "Y" if r["simo"] else "", r["tmu"],
                   lv.get("level") or "", role, cs, lv.get("parent_countersignature") or "",
                   lv.get("order") or "", lv.get("number") or "", lv.get("number_count") or ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _level_rows(data: dict[str, Any]) -> list[level_engine.LevelRow]:
    rows = []
    for r in data["rows"]:
        lv = r.get("level") or {}
        cyc = r.get("cycle") or {}
        rows.append(level_engine.LevelRow(
            content=(cyc.get("narrative") or r.get("sub_activity") or f"列{r['seq_no']}"),
            raw_seconds=float(lv.get("second") or 0), coefficient=1.0,
            number=lv.get("number") or None, number_count=lv.get("number_count"),
            ascription=lv.get("ascription") or None, level=lv.get("level") or None,
            countersignature=lv.get("countersignature") or None,
            parent_countersignature=lv.get("parent_countersignature") or None,
            order=lv.get("order"),
        ))
    return rows


async def to_lb_csv(session: AsyncSession, ws_id: uuid.UUID) -> str:
    """LB 可用的 CSV（Level System 欄位）。欄名可依 LB 規格再調。"""
    data = await read_worksheet(session, ws_id)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["seq_no", "content", "second", "ascription", "level", "level_min", "level_max", "variable",
                "countersignature", "parent_countersignature", "order",
                "number", "number_count", "machine_count", "manpower"])
    for r in data["rows"]:
        lv = r.get("level") or {}
        cyc = r.get("cycle") or {}
        vals = None
        if lv.get("level") not in (None, ""):
            vals, _ = level_engine._parse_level(str(lv["level"]))
        levels = sorted(vals) if vals else []
        w.writerow([r["seq_no"], cyc.get("narrative") or r.get("sub_activity") or f"列{r['seq_no']}",
                    lv.get("second"), lv.get("ascription"), lv.get("level"),
                    (levels[0] if levels else ""), (levels[-1] if levels else ""), (len(levels) > 1),
                    lv.get("countersignature"), lv.get("parent_countersignature"), lv.get("order"),
                    lv.get("number"), lv.get("number_count"), lv.get("machine_count"), lv.get("manpower")])
    return buf.getvalue()


async def lb_api_payload(session: AsyncSession, ws_id: uuid.UUID) -> dict[str, Any]:
    """LB API adapter（接口）：先回「將送出的」payload（build_output 合約）。

    LB request model 由 User 提供後，於此對映/POST 到 LB endpoint。
    """
    data = await read_worksheet(session, ws_id)
    output = level_engine.build_output(_level_rows(data))
    return {
        "status": "interface-ready",
        "note": "LB API 接口已備；待提供 LB request model 與 endpoint 後對映送出。目前為 dry-run。",
        "worksheet_id": str(ws_id),
        "payload": output,
    }


_STATUS_ZH = {"draft": "草稿", "approved": "已核准", "retired": "已退役"}


async def to_report_xlsx_bytes(session: AsyncSession, ws_id: uuid.UUID) -> bytes:
    """三 sheet 管理報表：案件資訊 / 動作明細 / 簽核歷程。"""
    from openpyxl import Workbook

    ws_model = await session.get(MostWorksheet, ws_id)
    if ws_model is None:
        raise WorksheetNotFound(str(ws_id))

    pv = await session.get(ProcessVersion, ws_model.process_version_id)
    sku = await session.get(Sku, pv.sku_id)
    product = await session.get(Product, sku.product_id)
    site = await session.get(Site, product.site_id)

    # 計算彙總；allowance_percent 為百分比值（e.g. 15.0 代表 15%），None = 未設定。
    # allowance 未設定時標準秒留空，不得以 normal 假充 standard。
    data = await read_worksheet(session, ws_id)
    total_tmu = data["total_tmu"] or 0.0
    wi_row_count = len(data["rows"])
    total_normal_sec = round(total_tmu * TMU_TO_SEC, 3)
    allowance_percent: float | None = data.get("allowance_percent")  # None or float %
    total_standard_sec: float | None = (
        round(total_normal_sec * (1 + allowance_percent / 100), 3)
        if allowance_percent is not None
        else None
    )

    process_name = sku.name_zh or sku.sku_code

    wb = Workbook()

    # ── Sheet 1：案件資訊（垂直 key-value）──
    ws1 = wb.active
    ws1.title = "案件資訊"
    rows_info = [
        ("廠區", site.name_zh),
        ("產品", product.name_zh),
        ("機種", sku.name_zh or sku.sku_code),
        ("製程", process_name),
        ("版本號", pv.version_no),
        ("狀態", _STATUS_ZH.get(pv.status, pv.status)),
        ("步驟總數", wi_row_count),
        ("總 TMU", round(total_tmu, 2)),
        ("總正常秒", total_normal_sec),
        ("寬放率", allowance_percent if allowance_percent is not None else ""),
        ("總標準秒", total_standard_sec if total_standard_sec is not None else ""),
        ("核准時間", pv.published_at.strftime("%Y-%m-%d %H:%M:%S") if pv.published_at else ""),
    ]
    for key, val in rows_info:
        ws1.append([key, val])

    # ── Sheet 2：動作明細──
    ws2 = wb.create_sheet("動作明細")
    ws2.append(["序號", "說明", "序列模型", "頻率", "插槽明細", "步驟TMU", "正常秒(4位)", "標準秒(4位)", "SIMO組"])
    for r in data["rows"]:
        cyc = r.get("cycle") or {}
        slot_inputs: dict[str, Any] = cyc.get("slot_inputs") or {}
        # 展開 slot_inputs：把各 key 的 value dict 中各欄的 key=value 連接
        slot_parts: list[str] = []
        for slot_key, slot_val in slot_inputs.items():
            if isinstance(slot_val, dict):
                for k, v in slot_val.items():
                    if k != "manual_override" and v not in (None, 0, "", [], {}):
                        slot_parts.append(f"{slot_key}({k})={v}")
            else:
                if slot_val not in (None, 0, "", [], {}):
                    slot_parts.append(f"{slot_key}={slot_val}")
        slot_detail = " | ".join(slot_parts) if slot_parts else ""

        step_tmu = float(cyc.get("total_tmu") or 0)
        freq = float(r.get("frequency") or 1)
        normal_sec = round(step_tmu * TMU_TO_SEC, 4)
        standard_sec: float | None = (
            round(normal_sec * (1 + allowance_percent / 100), 4)
            if allowance_percent is not None
            else None
        )

        simo_gid = r.get("simo_group_id") or ""
        simo_display = str(simo_gid)[:8] if simo_gid else ""

        ws2.append([
            r["seq_no"],
            r.get("sub_activity") or "",
            cyc.get("seq_kind") or "",
            freq,
            slot_detail,
            step_tmu,
            normal_sec,
            standard_sec if standard_sec is not None else "",
            simo_display,
        ])

    # ── Sheet 3：簽核歷程──
    ws3 = wb.create_sheet("簽核歷程")
    ws3.append(["時間", "動作", "從狀態", "至狀態", "執行者", "備註"])
    audit_rows = (
        await session.execute(
            select(WorkflowAuditLog)
            .where(
                WorkflowAuditLog.entity_type == "process_version",
                WorkflowAuditLog.entity_id == pv.id,
            )
            .order_by(WorkflowAuditLog.created_at.asc())
        )
    ).scalars().all()
    for entry in audit_rows:
        ws3.append([
            entry.created_at.strftime("%Y-%m-%d %H:%M:%S") if entry.created_at else "",
            entry.action,
            _STATUS_ZH.get(entry.from_status or "", entry.from_status or ""),
            _STATUS_ZH.get(entry.to_status or "", entry.to_status or ""),
            entry.actor,
            entry.comment or "",
        ])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
